# -*- coding: utf-8 -*-
"""法律研究模块测试：引用可解析且字段完整 / 多查询去重 / 诚实缺口 / DOCX 交付物。"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402

from app import research, research_report  # noqa: E402
from app.corpus import LawCorpus  # noqa: E402

C = LawCorpus()

CARD_FIELDS = (
    "law_id", "law_title", "law_status", "effective_date", "promulgation_instrument",
    "article_no", "article_label", "chapter", "text", "source_url", "source_kind", "score",
)


def test_memo_cards_resolvable_and_complete():
    memo = research.build_research_memo("民间借贷的利率上限")
    assert memo["cards"], "应有命中"
    for card in memo["cards"]:
        cit = C.citation_of(card["law_id"], card["article_no"])  # 引用不存在会抛 KeyError
        assert cit["law_title"] == card["law_title"]
        assert cit["text"] == card["text"]
        for f in CARD_FIELDS:
            assert f in card, f"卡片缺字段：{f}"
    assert memo["meta"]["method"] in {"bm25-controlled-terms", "bm25-controlled-groups"}
    assert memo["meta"]["corpus_size"] == len(C.articles)
    assert memo["disclaimer"]
    # references 与 cards 一一对应，全部可溯源
    assert len(memo["references"]) == len(memo["cards"])
    for r in memo["references"]:
        assert r["kind"] == "statute" and r["law_title"] and r["source_url"] and r["text"]


def test_memo_no_duplicate_cards_and_sorted():
    memo = research.build_research_memo("网购七日无理由退货与逾期利息", top_k=12)
    keys = [(c["law_id"], c["article_no"]) for c in memo["cards"]]
    assert len(keys) == len(set(keys)), "多查询合并不得产生重复卡"
    scores = [c["score"] for c in memo["cards"]]
    assert scores == sorted(scores, reverse=True)
    assert len(memo["cards"]) <= 12


def test_plain_language_consumer_query_prioritizes_consumer_sources():
    """自然问句的套话不得压过实体主题；“网购假货”首批必须来自消费者/电商语料。"""
    memo = research.build_research_memo("网购到假货可以核对哪些现行法条", top_k=8)
    assert memo["cards"]
    assert memo["cards"][0]["law_id"] in {"cl-2013", "crpl-imp-2024", "ecom-2018", "wlxf-2022"}
    assert not any(c["law_id"] == "pcl-2023" for c in memo["cards"][:4])
    assert len({c["law_id"] for c in memo["cards"][:4]}) >= 2, "不得被一次编章扩展垄断首批结果"
    assert {"消费者", "经营者", "欺诈"} <= set(memo["issue_frame"]["keywords"])


def test_memo_issue_frame_is_programmatic():
    memo = research.build_research_memo("民间借贷的利率上限是多少？")
    frame = memo["issue_frame"]
    assert "民间借贷的利率上限是多少？" in frame["restate"]
    assert "依据检索" in frame["restate"]
    assert frame["keywords"], "应提取到关键词"
    single_stops = {w for w in research.STOPWORDS if len(w) == 1}
    for kw in frame["keywords"]:
        assert kw not in research.STOPWORDS
        assert not any(ch in single_stops for ch in kw), f"关键词混入停用词字：{kw}"
    # framework 与 cards 聚类一致
    hit_total = sum(f["hit_count"] for f in memo["framework"])
    assert hit_total == len(memo["cards"])
    for f in memo["framework"]:
        assert isinstance(f["chapters"], list)


def test_memo_gaps_on_irrelevant_question():
    memo = research.build_research_memo("asdfgh qwerty")
    assert memo["gaps"], "无关问题应有诚实缺口"
    assert len(memo["cards"]) <= 2
    assert ("未检索到" in memo["gaps"][0]) or ("不作推断" in memo["gaps"][0])
    assert memo["references"] == []


def test_memo_normal_question_hits_civil_code():
    memo = research.build_research_memo("民间借贷的利率上限")
    assert memo["cards"]
    assert any(c["law_id"] == "civl-2020" for c in memo["cards"])
    assert memo["framework"]
    assert memo["gaps"] == [], "命中充足时不应有缺口提示"


def test_memo_law_ids_scope_filter():
    memo = research.build_research_memo("商品质量不合格怎么退货", law_ids=["cl-2013"])
    assert memo["cards"], "限定单法也应能命中"
    assert all(c["law_id"] == "cl-2013" for c in memo["cards"])
    assert memo["scope"]["law_ids"] == ["cl-2013"]
    try:
        research.build_research_memo("任意问题", law_ids=["not-a-law"])
    except ValueError:
        pass
    else:
        raise AssertionError("未知法律编号应抛 ValueError")


def test_report_docx_generated():
    memo = research.build_research_memo("民间借贷的利率上限")
    data = research_report.generate_research_docx(memo)
    assert len(data) > 5000
    assert data[:2] == b"PK", "DOCX 应为 zip 包（PK 魔数）"
    doc = Document(io.BytesIO(data))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "法律研究备忘录" in joined
    assert "参考依据" in joined
    assert "免责声明" in joined
    # 表格：法律框架表 + 参考依据表
    assert len(doc.tables) == 2
    ref_table = doc.tables[1]
    assert ref_table.rows[0].cells[0].text == "编号"
    assert len(ref_table.rows) == 1 + len(memo["cards"])


def test_report_docx_with_empty_hits_is_honest():
    memo = research.build_research_memo("asdfgh qwerty")
    data = research_report.generate_research_docx(memo)
    assert len(data) > 5000
    doc = Document(io.BytesIO(data))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "未命中任何条文" in joined
    assert "不作任何推断" in joined


def test_chapter_expansion_tolerates_none(monkeypatch):
    """A5 回归：第三组查询用首条命中 chapter 扩展，chapter=None 时不得崩溃（语料 schema 允许 None）。"""
    from app import research
    from app.corpus import get_corpus as real
    base = real()
    class FakeCorpus:
        def __getattr__(self, name):
            return getattr(base, name)
        def search(self, q, top_k=8, law_id=None):
            hits = base.search(q, top_k=top_k, law_id=law_id)
            if hits:
                hits = [dict(hits[0], chapter=None)] + hits[1:]
            return hits
    monkeypatch.setattr(research, "get_corpus", lambda: FakeCorpus())
    memo = research.build_research_memo("网络购物七日无理由退货")
    assert memo["cards"]


def test_consumer_question_is_consistent_across_public_retrieval_paths():
    """同一公众问题在搜索、引用式问答、研究和事实计划中必须走同一受控主题编排。"""
    from fastapi.testclient import TestClient
    from app.main import app
    from app.retrieval_terms import orchestrated_search

    query = "网购到假货可以核对哪些现行法条"

    def assert_consumer(cards, no_key="article_no"):
        keys = {(card["law_id"], card[no_key]) for card in cards}
        assert ("cl-2013", 55) in keys
        assert any(key in keys for key in {
            ("cl-2013", 24), ("cl-2013", 25), ("cl-2013", 44), ("ecom-2018", 38),
        })
        assert not any(card["law_id"] == "pcl-2023" for card in cards[:4])

    hits, meta = orchestrated_search(C, query, top_k=8)
    assert meta["method"] == "bm25-controlled-groups"
    assert_consumer(hits, "no")

    client = TestClient(app)
    search = client.get("/api/search", params={"q": query, "top_k": 8}).json()
    assert_consumer(search["hits"], "no")
    assert search["retrieval_meta"]["matched_groups"] == [
        "consumer-fraud", "consumer-return", "online-platform",
    ]

    answer = client.post("/api/qa/ask", json={"question": query, "top_k": 8}).json()
    assert_consumer(answer["answer_cards"])

    memo = client.post("/api/research/memo", json={"question": query, "top_k": 8}).json()
    assert_consumer(memo["cards"])

    plan = client.post("/api/needs/plan", json={
        "summary": "我在网上商店买到假货，商家拒绝退货",
        "trigger": "网购商品后发现是假货",
        "timeline": ["收到商品后核对发现与宣传不符"],
        "actual_outcome": "商家拒绝退款",
        "parties": ["消费者", "网络商家"],
        "evidence_owned": ["订单与聊天记录"],
        "desired_outcome": "了解退货退款路径",
    }).json()
    assert_consumer(plan["articles"])


def test_unbound_scenario_guidance_is_not_publicly_served():
    from fastapi.testclient import TestClient
    from app.main import app
    client = TestClient(app)
    assert client.get("/api/scenarios").status_code == 410
    assert client.get("/api/scenarios/match", params={"text": "欠薪"}).status_code == 410
