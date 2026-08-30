# -*- coding: utf-8 -*-
"""法律研究备忘录 DOCX 渲染（python-docx）。

数据纪律：文档中的每一条依据都来自入参 memo（其引用已由 research 模块经
corpus.citation_of 逐条校验），语料快照信息只读复用 drafting._snapshot()
的日期与 manifest 字段——报告层不补充任何语料之外的事实，无命中时如实留空。
排版沿用 docxgen 约定：A4、公文页边距（上3.7/下3.5/左2.8/右2.6 cm）、
仿宋正文、黑体标题；页脚固定「LegalHigh 原型生成 · 内容不构成法律意见」。
"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from .docxgen import SMALL_SIZE, _para, _set_font
from .drafting import _snapshot

_TABLE_SIZE = Pt(10.5)  # 五号：表格与来源附注
_GRAY = RGBColor(0x63, 0x63, 0x66)


def _heading(doc, text):
    return _para(doc, text, indent=False, size=Pt(14), east="黑体", bold=True,
                 line=Pt(26), before=8, after=4)


def _cell_text(cell, text, *, bold=False, east="仿宋", size=_TABLE_SIZE):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = size
    run.font.bold = bold
    _set_font(run, east)


def _table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        _cell_text(t.rows[0].cells[j], h, bold=True, east="黑体")
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            _cell_text(t.rows[i].cells[j], value)
    return t


def _footer_note(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LegalHigh 原型生成 · 内容不构成法律意见")
    run.font.size = SMALL_SIZE
    run.font.color.rgb = RGBColor(0x8E, 0x8E, 0x93)
    _set_font(run, "宋体")


def generate_research_docx(memo: dict) -> bytes:
    """把 build_research_memo 的结果渲染为可交付的 DOCX（bytes）。"""
    snap = _snapshot()  # 只读快照：生成日期 + 语料 manifest（不改写、不落库）
    manifest = snap.get("corpus_manifest", {})
    meta = memo.get("meta", {})
    scope = memo.get("scope", {})
    frame = memo.get("issue_frame", {})
    laws_meta = manifest.get("laws", [])
    article_total = meta.get("corpus_size") or sum(l.get("article_count", 0) for l in laws_meta)

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(3.7), Cm(3.5)
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.6)

    # 标题与快照信息
    _para(doc, "法律研究备忘录", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
          size=Pt(18), east="黑体", bold=True, line=Pt(34), after=8)
    _para(
        doc,
        f"生成日期：{snap.get('generated_at', '?')}　|　语料快照：{len(laws_meta)} 部法规 / "
        f"{article_total} 条 · 采集日期 {manifest.get('fetch_date') or '?'}　|　"
        f"检索方法：{meta.get('method', '?')}",
        indent=False, size=Pt(10), east="宋体", line=Pt(18), after=10)

    # 一、问题界定（程序化改写，不含观点）
    _heading(doc, "一、问题界定")
    _para(doc, f"研究问题：{memo.get('question', '')}")
    _para(doc, f"问题改写（程序化，不代表法律评价）：{frame.get('restate', '')}")
    keywords = frame.get("keywords") or []
    _para(doc, "检索关键词：" + ("、".join(keywords) if keywords else "（未提取到有效关键词）"))
    scope_ids = scope.get("law_ids")
    scope_desc = f"限定于 {len(scope_ids)} 部法律（{'、'.join(scope_ids)}）" if scope_ids else "全部语料"
    _para(doc, f"检索范围：{scope_desc}；多组查询合并去重后按相关度取前 {scope.get('top_k', 12)} 条。")

    # 二、法律框架（编章聚类表）
    _heading(doc, "二、法律框架（命中条文的编章聚类）")
    fw = memo.get("framework") or []
    if fw:
        _table(doc, ["法律", "编·章", "命中数"], [
            [f"《{f['law_title']}》（{f['law_id']}）",
             "；".join(f.get("chapters") or []),
             str(f.get("hit_count", 0))]
            for f in fw
        ])
    else:
        _para(doc, "本次检索未命中任何条文，无法构建法律框架。")

    # 三、相关条文（完整文本 + 版本/施行日期/来源）
    _heading(doc, "三、相关条文（完整文本）")
    cards = memo.get("cards") or []
    if not cards:
        _para(doc, "无命中条文——本备忘录不包含任何推断性内容。")
    for i, c in enumerate(cards, 1):
        _para(
            doc,
            f"{i}. 《{c['law_title']}》{c['article_label']}"
            f"（{c['law_status']}，施行日期：{c['effective_date'] or '见官方文本'}；编章：{c['chapter']}）",
            indent=False, size=Pt(12), east="黑体", line=Pt(24), before=6)
        for line in (c.get("text") or "").split("\n"):
            if line.strip():
                _para(doc, line)
        _para(doc, f"来源：{c['source_url']}（{c['source_kind']}）", indent=False,
              size=SMALL_SIZE, east="宋体", color=_GRAY, line=Pt(16))

    # 四、检索说明与缺口（如实列出，不粉饰）
    _heading(doc, "四、检索说明与缺口")
    _para(
        doc,
        f"检索方法：{meta.get('method', '?')}——对问题做多组查询（原句 / 去停用词关键词 / "
        f"依首次命中条文编章标题扩展一次），结果按（法律编号, 条号）合并去重、得分取最大值后排序；"
        f"语料规模 {article_total} 条。")
    _para(doc, "实际执行查询：", indent=False)
    for i, q in enumerate(meta.get("queries") or [], 1):
        _para(doc, f"（{i}）{q}", indent=False)
    gaps = memo.get("gaps") or []
    if gaps:
        for g in gaps:
            _para(doc, f"缺口提示：{g}")
    else:
        _para(doc, "本次检索命中依据数量充足，未触发缺口提示。")

    # 附录：参考依据表
    _heading(doc, "附录：参考依据表")
    refs = memo.get("references") or []
    if refs:
        _table(doc, ["编号", "类型", "出处", "链接"], [
            [str(i), "法条（statute）",
             f"《{r['law_title']}》{r['article_label']}（{r.get('status', '')}）",
             r.get("source_url", "")]
            for i, r in enumerate(refs, 1)
        ])
    else:
        _para(doc, "无参考依据。")

    _para(doc, "免责声明：" + memo.get("disclaimer", ""), indent=False, size=Pt(10),
          east="宋体", color=_GRAY, line=Pt(18), before=12)

    _footer_note(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
