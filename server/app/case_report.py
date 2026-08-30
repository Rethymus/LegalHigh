# -*- coding: utf-8 -*-
"""案件分析报告 DOCX 渲染（python-docx，内存生成，不落盘）。

数据纪律：文档中的每一条内容都来自入参 analysis（其 statute 引用已由 legalmodel
经 corpus.citation_of 逐条校验，文本片段均为用户文本的原文切片）——报告层不补充
任何入参之外的事实。顶部固定红色横幅「分析草稿 · 非法律意见 · 不作心理诊断」；
行为模式章节必须完整转载固定非诊断声明（心理红线）。
排版沿用 docxgen 约定：A4、公文页边距（上3.7/下3.5/左2.8/右2.6 cm）、仿宋正文、
黑体标题；页脚「LegalHigh 原型生成 · 内容不构成法律意见」。
"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from .docxgen import SMALL_SIZE, _banner, _para, _set_font

_TABLE_SIZE = Pt(10.5)  # 五号：表格
_SMALL_SIZE = Pt(9)
_GRAY = RGBColor(0x63, 0x63, 0x66)

_STATUS_LABEL = {
    "supported": "有文本支持（supported）",
    "unverified": "待补充证据或线索（unverified）",
}
_LEVEL_LABEL = {"high": "高", "medium": "中", "low": "低", "absent": "未出现"}


def _heading(doc, text):
    return _para(doc, text, indent=False, size=Pt(14), east="黑体", bold=True,
                 line=Pt(26), before=8, after=4)


def _cell_text(cell, text, *, bold=False, east="仿宋", size=_TABLE_SIZE):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = size
    run.font.bold = bold
    _set_font(run, east)


def _table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        _cell_text(t.rows[0].cells[j], h, bold=True, east="黑体")
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            _cell_text(t.rows[i].cells[j], value)
    return t


def _excerpt_line(doc, prefix: str, excerpt: str):
    _para(doc, f"{prefix}「{excerpt}」", indent=False, size=Pt(12), east="宋体",
          line=Pt(22), after=2)


def _footer_note(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LegalHigh 原型生成 · 内容不构成法律意见")
    run.font.size = SMALL_SIZE
    run.font.color.rgb = RGBColor(0x8E, 0x8E, 0x93)
    _set_font(run, "宋体")


def generate_case_docx(analysis: dict) -> bytes:
    """把 case_analysis.analyze_case 的结果渲染为可交付的 DOCX（bytes，仅内存）。"""
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(3.7), Cm(3.5)
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.6)

    claim = analysis.get("claim") or {}
    profile = analysis.get("profile") or {}
    behavior = analysis.get("behavior") or {}
    summary = analysis.get("summary") or {}

    # 标题 + 顶部红字横幅（草稿属性声明）
    _para(doc, "案件分析报告（草稿）", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
          size=Pt(18), east="黑体", bold=True, line=Pt(34), after=4)
    _banner(doc, "分析草稿 · 非法律意见 · 不作心理诊断")

    # 案件概述（仅程序化字段）
    _para(doc,
          f"案件名称：{analysis.get('title', '未命名案件')}　|　生成日期："
          f"{analysis.get('generated_at', '?')}　|　请求权类型："
          f"{(claim.get('claim') or {}).get('name', '')}",
          indent=False, size=Pt(10), east="宋体", line=Pt(18), after=8)

    # 一、当事人人像（每方：角色/线索/行为摘录；时间线为全局线索）
    _heading(doc, "一、当事人人像")
    parties = profile.get("parties") or []
    if not parties:
        _para(doc, "所提供文本中未抽到当事人角色线索（确定性抽取，不推测）。")
    for p in parties:
        hint = p.get("name_hint")
        _para(doc, f"【{p['role']}】" + (f"名称线索：{hint}" if hint else "名称线索：（文本中未见）"),
              indent=False, size=Pt(12), east="黑体", line=Pt(24), before=4)
        for m in p.get("mentions") or []:
            _excerpt_line(doc, "出现：", m["excerpt"])
        for b in p.get("behaviors") or []:
            _excerpt_line(doc, f"{b['type']}类：", b["excerpt"])
    timeline = profile.get("timeline") or []
    if timeline:
        _para(doc, "时间线（全局日期线索，均摘自所提供文本）：", indent=False,
              size=Pt(12), east="黑体", line=Pt(24), before=4)
        for t in timeline:
            _excerpt_line(doc, f"{t['date_hint']}：", t["excerpt"])
    else:
        _para(doc, "时间线：所提供文本中未抽到日期线索。")

    # 二、沟通行为模式分析（指标表 + 固定非诊断声明全文）
    _heading(doc, "二、沟通行为模式分析")
    indicators = behavior.get("indicators") or []
    _table(doc, ["指标", "等级", "依据片段（原文摘录）"], [
        [f"{i['id']} {i['title']}", _LEVEL_LABEL.get(i["level"], i["level"]),
         "　".join(s["excerpt"] for s in (i.get("spans") or [])[:2]) or "（未命中）"]
        for i in indicators
    ])
    for i in indicators:
        if i["level"] != "absent":
            _para(doc, f"{i['id']} {i['title']}：{i['note']}", indent=False,
                  size=Pt(11), east="宋体", line=Pt(20), before=4)
            _para(doc, f"　应对建议：{i['advice']}", indent=False, size=Pt(11),
                  east="宋体", line=Pt(20))
    _para(doc, "【固定声明】" + behavior.get("fixed_disclaimer", ""), indent=False,
          size=Pt(11), east="黑体", color=RGBColor(0x8A, 0x6D, 0x00), line=Pt(20), before=6)

    # 三、法律要件矩阵（状态/证据片段/依据条文全文含施行日期）
    _heading(doc, "三、法律要件矩阵")
    elements = claim.get("elements") or []
    if not elements:
        _para(doc, "无要件模型匹配结果。")
    for el in elements:
        _para(doc, f"{el['id']} {el['title']}——{ _STATUS_LABEL.get(el['status'], el['status']) }",
              indent=False, size=Pt(12), east="黑体", line=Pt(24), before=6)
        if el.get("evidence_spans"):
            for sp in el["evidence_spans"]:
                _excerpt_line(doc, "证据片段：", sp["excerpt"])
        else:
            _para(doc, "证据片段：（所提供文本中未见对应线索，待补充证据）", indent=False,
                  size=Pt(12), east="宋体", line=Pt(22))
        for cit in el.get("citations") or []:
            eff = cit.get("effective_date") or "见官方文本"
            _para(doc, f"依据条文：《{cit['law_title']}》{cit['article_label']}"
                       f"（{cit.get('status', '')}，施行日期：{eff}）",
                  indent=False, size=Pt(11), east="宋体", line=Pt(20))
            for line in (cit.get("text") or "").split("\n"):
                if line.strip():
                    _para(doc, line)
            _para(doc, f"来源：{cit.get('source_url', '')}", indent=False,
                  size=_SMALL_SIZE, east="宋体", color=_GRAY, line=Pt(16))

    # 四、综合提示（程序化 summary，不加观点）
    _heading(doc, "四、综合提示")
    _para(doc,
          f"本案文本中确定性识别到 {summary.get('profile_parties', 0)} 方当事人角色；"
          f"行为模式指标 {summary.get('behavior_active', 0)}/6 项出现命中；"
          f"请求权「{(claim.get('claim') or {}).get('name', '')}」"
          f"{summary.get('claim_supported', 0)} 项要件有文本支持、"
          f"{summary.get('claim_unverified', 0)} 项待补充证据或线索"
          "（关键词级匹配，不构成法律意见）。")
    _para(doc, "要件矩阵：" + (claim.get("summary") or {}).get("overall", ""))

    # 附录：参考依据表（statute 带 URL；text_span 计数）
    _heading(doc, "附录：参考依据表")
    refs = analysis.get("references") or []
    if refs:
        rows = []
        for i, r in enumerate(refs, 1):
            if r.get("kind") == "statute":
                rows.append([str(i), "法条（statute）",
                             f"《{r.get('law_title', '')}》{r.get('article_label', '')}"
                             f"（{r.get('status', '')}）",
                             r.get("source_url", "")])
            else:
                rows.append([str(i), "文本片段（text_span）",
                             f"{r.get('label', '')} · 命中 {r.get('count', 0)} 处（均取自所提供文本）",
                             "—"])
        _table(doc, ["编号", "类型", "出处", "链接"], rows)
    else:
        _para(doc, "无参考依据。")

    # 免责声明（各模块声明合并转载）
    _heading(doc, "免责声明")
    for d in analysis.get("disclaimers") or []:
        _para(doc, "· " + d, indent=False, size=Pt(10), east="宋体", color=_GRAY, line=Pt(18))

    _footer_note(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
