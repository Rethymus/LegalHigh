# -*- coding: utf-8 -*-
"""DOCX 排版渲染（python-docx，MIT）：让「生成即可交付」落地。

排版取实务惯例参数：A4、公文式页边距（上3.7/下3.5/左2.8/右2.6 cm）、正文仿宋四号
（14pt）、行距固定 28pt、首行缩进两字符、标题黑体居中；引用条文随文附版本与施行日期。
合规护栏：律师函在「律师签发」前输出红色草稿横幅；未核验文书输出提示横幅。
"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_SIZE = Pt(14)
TITLE_SIZE = Pt(18)
HEADING_SIZE = Pt(14)
SMALL_SIZE = Pt(9)


def _set_font(run, east_asia: str, ascii_font: str = "Times New Roman"):
    run.font.name = ascii_font
    r = run._element.rPr.rFonts
    r.set(qn("w:eastAsia"), east_asia)


def _para(doc, text, *, align=None, indent=True, size=BODY_SIZE, east="仿宋", bold=False,
          color=None, line=Pt(28), before=0, after=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pf.first_line_indent = Pt(size.pt * 2)
    if align is not None:
        pf.alignment = align
    run = p.add_run(text)
    run.font.size = size
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    _set_font(run, east)
    return p


def _banner(doc, text, color=RGBColor(0xD7, 0x00, 0x00)):
    _para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=Pt(12),
          east="黑体", bold=True, color=color, line=Pt(20), after=6)


def _footer_note(doc, draft):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    snap = draft.get("snapshot") or {}
    gen = snap.get("generated_at", "?")
    note = f"LegalHigh 原型生成 · 文书编号 {draft.get('id', '')} · 生成日期 {gen} · 法条版本快照 {snap.get('corpus_manifest', {}).get('fetch_date', '?')} · 内容不构成法律意见"
    run = p.add_run(note)
    run.font.size = SMALL_SIZE
    run.font.color.rgb = RGBColor(0x8E, 0x8E, 0x93)
    _set_font(run, "宋体")


def generate_docx(draft: dict) -> bytes:
    """draft 为 storage.get_draft() 的返回（含 status / content / snapshot）。"""
    content = draft["content"]
    sections = content["sections"]
    gate_note = content.get("gate_note")
    status = draft.get("status", "draft")
    template_id = draft.get("template_id")

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(3.7), Cm(3.5)
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.6)

    # 合规横幅：状态机驱动的可交付性标注
    if template_id == "lawyer_letter" and status != "issued":
        _banner(doc, "草稿 · 未经执业律师核验签发 —— 不得以律所/律师名义对外发送（《律师法》第13条）")
    elif template_id != "lawyer_letter" and status == "draft":
        _banner(doc, "草稿 · 供内部审阅，未经人工核验定稿", color=RGBColor(0x8A, 0x6D, 0x00))
    if status in ("verified", "issued"):
        meta = draft.get("verified_by") or ""
        _banner(doc, f"已通过人工核验（核验人：{meta}，{draft.get('verified_role') or ''}）",
                color=RGBColor(0x1F, 0x7A, 0x33))
        if status == "issued":
            _banner(doc, f"已签发（签发人：{draft.get('issued_by', '')}，{draft.get('issued_at', '')[:10]}）",
                    color=RGBColor(0x1F, 0x7A, 0x33))

    for s in sections:
        t = s["type"]
        if t == "title":
            _para(doc, s["text"], align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=TITLE_SIZE,
                  east="黑体", bold=True, line=Pt(34), after=8)
        elif t == "subtitle":
            _para(doc, s["text"], align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=Pt(12),
                  east="宋体", line=Pt(22), after=12)
        elif t == "heading":
            _para(doc, s["text"], indent=False, size=HEADING_SIZE, east="黑体", bold=True,
                  line=Pt(26), before=8, after=4)
        elif t == "numbered":
            _para(doc, f"{s['n']}、{s['text']}")
        elif t == "party":
            for line in s["lines"]:
                _para(doc, line, indent=False)
        elif t == "closing":
            for i, line in enumerate(s["text"].split("\n")):
                _para(doc, line, indent=False, before=8 if i == 0 else 0)
        elif t == "signature":
            for line in s["lines"]:
                _para(doc, line, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False, before=6)
        else:  # para / para_noindent
            _para(doc, s["text"], indent=(t == "para"))

    # 草稿说明仅对未签发文书显示；已签发版本是对外交付物，不再携带草稿警示
    if gate_note and status != "issued":
        _para(doc, "【说明】" + gate_note, indent=False, size=Pt(10), east="宋体",
              color=RGBColor(0x63, 0x63, 0x66), line=Pt(18), before=10)

    _footer_note(doc, draft)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- 审查记录 DOCX（Word 修订双轨·M7-T1 前半） ----------------

def _add_tracked_insert(paragraph, text: str, author: str, date: str, doc):
    """在段落中追加一段「修订插入」文本（w:ins 包裹 w:r）——Word/WPS 打开即为修订标记，
    律师可接受/拒绝。python-docx 无原生 API，直接操作底层 XML。"""
    from docx.oxml.ns import qn
    run = paragraph.add_run(text)
    r_el = run._r
    ins = r_el.makeelement(qn("w:ins"), {})
    ins.set(qn("w:id"), str(doc._next_id))
    doc._next_id += 1
    ins.set(qn("w:author"), author)
    ins.set(qn("w:date"), date)
    r_el.addprevious(ins)
    ins.append(r_el)
    return run


def generate_review_docx(review: dict) -> bytes:
    """审查记录导出：合同条款原文 + 每条 AI 建议以「修订插入」写入（作者=LegalHigh AI）。
    修订版式由结构化数据（findings.suggestion）决定，不让自由生成决定版式。"""
    import re
    from datetime import datetime, timezone
    from docx import Document
    doc = Document()
    doc._next_id = 1
    date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    doc.add_heading(review.get("title") or "合同审查记录", level=0)
    doc.add_paragraph(f"审查时间：{review.get('created_at', '')}　审查点：{review['result']['engine_meta']['checkpoint_count']} 个　风险：高 {review['result']['summary']['high']} / 中 {review['result']['summary']['medium']} / 低 {review['result']['summary']['low']}")
    by_clause: dict = {}
    for f in review["result"]["findings"]:
        by_clause.setdefault(f.get("clause_id") or "__whole__", []).append(f)
    def _bookmark(paragraph, name: str):
        """书签锚定回传解析的定位（M7-T1 后半）：LH_{finding_id} 包裹建议段落。"""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        bs = OxmlElement("w:bookmarkStart")
        bs.set(qn("w:id"), str(doc._next_id))
        bs.set(qn("w:name"), name)
        be = OxmlElement("w:bookmarkEnd")
        be.set(qn("w:id"), str(doc._next_id))
        doc._next_id += 1
        paragraph._p.addprevious(bs)
        paragraph._p.addnext(be)

    for c in review["result"]["clauses"]:
        doc.add_heading(c.get("heading") or c["label"], level=2)
        p = doc.add_paragraph(c["text"])
        for f in by_clause.get(c["id"], []):
            sug = f.get("suggestion", "").strip()
            if not sug:
                continue
            p2 = doc.add_paragraph()
            _bookmark(p2, f"LH_{f['id']}")
            p2.add_run(f"[{f['checkpoint_title']}｜{ {'high': '高风险', 'medium': '中风险', 'low': '低风险' }[f['risk']]}] ")
            _add_tracked_insert(p2, f"建议：{sug}", "LegalHigh AI", date, doc)
    for f in by_clause.get("__whole__", []):
        p3 = doc.add_paragraph()
        _bookmark(p3, f"LH_{f['id']}")
        p3.add_run(f"[全文级｜{f['checkpoint_title']}] ")
        _add_tracked_insert(p3, f"建议：{f['suggestion']}", "LegalHigh AI", date, doc)
    doc.add_paragraph(review["result"]["disclaimer"])

    from io import BytesIO
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
